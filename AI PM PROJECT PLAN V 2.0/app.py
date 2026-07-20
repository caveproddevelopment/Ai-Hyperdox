"""
AIPM - Planning Agent (v2.2, production wrapper restored)
Agent 2 of the AI Hyperdox pipeline: Goals & Scope -> Planning -> Execution.

WHAT THIS PASS CHANGES (and why):

The v2.2 agent logic (combined Guardrail PDF, validate_coverage(), validate_
unspecified_markers(), the Agent-1-matching input model) came in as a bare Gradio
app writing PDFs/JSON to /tmp only - the same situation Agent 1's v5.0 was received
in. That would have shipped without the Flask REST API, CORS, Firebase Storage
upload layer, or health check server that aihyperdox.com and Railway actually depend
on. This pass does two things, both mechanical, neither touching agent behavior:

1. RESTORES THE PRODUCTION WRAPPER, unchanged from how it works in this agent's own
   last production release (v2.0): Flask app, origin-restricted CORS (aihyperdox.com,
   the Vercel preview domain, and localhost:5173 for local frontend dev - NOT a
   blanket CORS(app) the way Agent 1 uses, since that's the policy this agent already
   shipped with), a port-8081 health-check thread, /api/predict, /download, and
   /api/health. See upload_to_storage() for the permanent-URL mechanism.

2. ADDS BASE64 FILE-UPLOAD SUPPORT, matching the fix already applied to Agent 1
   (Goals & Scope). v2.2's extract_text_from_file() and extract_handoff_json() only
   understood Gradio-native shapes (a plain string path, or an object with .name) -
   fine for the manual-test Blocks UI, but the React frontend's /api/predict call
   can't hand over a filesystem path or a live File object across JSON; it has to
   base64-encode the file via FileReader.readAsDataURL() and send a
   {"name": str, "data": str} dict instead (optionally still carrying a
   "data:<mime>;base64," prefix, stripped below). Both functions now accept that
   shape alongside the existing ones, and MAX_UPLOAD_FILE_BYTES is a new server-side
   size backstop independent of whatever the frontend enforces client-side. Every
   generated output (4 PDFs + the new Combined PDF + the Backbone Plan JSON) is now
   also run through upload_to_storage() before being returned, so ProjectPlanning.jsx
   gets the same permanent-URL {"url":..., "path":...} shape it already expects from
   Agent 1 - the bare v2.2 code was still returning raw /tmp paths, which don't
   survive a redeploy and can't be fetched directly by the browser.

WHAT v2.2 CHANGED FROM v2.1 (carried through unchanged in this pass):

1. COMBINED "PROJECT PLAN GUARDRAIL DOCUMENT" - matches Agent 1's pattern.
   v2.1 (and v2.0) produced 4 separate PDFs with no single document showing
   open_questions alongside the plan. Added a combined PDF (all 4 sections + Open
   Questions) mirroring Agent 1's "Scope Guardrail Document (Combined)" exactly.

2. COVERAGE GUARDRAIL - same class of fix as Agent 1's validate_coverage().
   validate_coverage() (phase-level) confirms every WBS phase has at least one
   Timeline row, one Resource Allocation row, and one Cost Baseline row referencing
   it, appending a specific open_questions entry for any gap.
   validate_unspecified_markers() is a second, narrower guardrail: if any field was
   marked UNSPECIFIED but open_questions came back empty, it flags that explicitly -
   the direct fix for an observed live bug where a Cost Management Plan came back
   100% "UNSPECIFIED - see open questions" with zero open_questions explaining why.

WHAT v2.1 CHANGED FROM v2.0 (carried through unchanged in this pass, input/UI only):

1. INPUT MODEL NOW MATCHES AGENT 1 (Goals & Scope). Replaced the single-purpose
   "Budget Spreadsheet" file upload (XLSX/CSV/PDF/TXT) with the same pattern as
   Agent 1: a pasted-text box for budget/financial data PLUS a multi-file "Upload
   Supporting Documents" field accepting any file type/count. Both optional and
   additive to the manual Milestones/Timeline/Resources fields.
   NOTE - HANDOFF SCHEMA CHANGE FOR THE FRONTEND: generate_project_plan()'s
   positional signature changed from v2.0's
     (project_name, milestones, timeline_input, resources, budget_file,
      methodology, handoff_upload)                                    [7 args]
   to v2.2's
     (project_name, milestones, timeline_input, resources, budget_text_input,
      supporting_uploads, methodology, handoff_upload)                [8 args]
   - budget_file (single spreadsheet) is gone; budget_text_input (str) and
   supporting_uploads (list) are new and sit in different positions. The old
   ProjectPlanning.jsx sends 6 positional items with budget hardcoded to null and
   never wires up a handoff upload at all - flag this to Aatish explicitly, the
   frontend needs a matching rewrite, not just a redeploy of this file.

2. UI LAYOUT NOW MATCHES AGENT 1's STRUCTURE EXACTLY. No custom CSS block.

WHAT v2.0 CHANGED FROM v1.0 (carried through unchanged in this pass):

1. REAL SYSTEM PROMPT (see agent_prompt.md for the readable reference copy).
2. INGESTS THE GOALS & SCOPE HANDOFF PACKAGE directly - manual fields are still
   available and are treated as overrides/additions.
3. ONE STRUCTURED JSON CALL instead of 4 isolated ones.
4. ANTI-GENERICISM GUARDRAILS + open_questions, mirroring Agent 1.
5. PRODUCES ITS OWN HANDOFF PACKAGE ("Backbone Plan" JSON) for the future Execution
   agent (Agent 3, not yet built).
6. PRODUCTION HARDENING: session-safe file naming, retries with backoff, input
   truncation, temperature tuned down, errors surfaced instead of raised.

NOTE ON MIGRATION: the only OpenAI-specific code lives in call_llm() - identical
shape to Agent 1's, so both can be swapped to Anthropic in the same pass later.
"""

import base64
import json
import os
import threading
import time
import urllib.parse
import uuid
from datetime import datetime, timezone
from http.server import HTTPServer, BaseHTTPRequestHandler
from io import BytesIO

import docx
import gradio as gr
from flask import Flask, request, jsonify, send_file, abort
from flask_cors import CORS
from openai import OpenAI
from PyPDF2 import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem,
    PageBreak,
)

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

# ── Firebase Admin ──────────────────────────────────────────────
import firebase_admin
from firebase_admin import credentials, storage as fb_storage

_firebase_ready = False


def init_firebase():
    global _firebase_ready
    if _firebase_ready:
        return True
    try:
        sa_json = os.getenv("FIREBASE_SERVICE_ACCOUNT")   # full JSON string
        bucket  = os.getenv("FIREBASE_STORAGE_BUCKET")    # e.g. your-project.appspot.com
        if not sa_json or not bucket:
            print("⚠ Firebase env vars not set — files will use /tmp fallback.")
            return False
        cred = credentials.Certificate(json.loads(sa_json))
        firebase_admin.initialize_app(cred, {"storageBucket": bucket})
        _firebase_ready = True
        print("✅ Firebase Admin initialised.")
        return True
    except Exception as e:
        print(f"⚠ Firebase init failed: {e}")
        return False


def upload_to_storage(local_path: str, project_name: str, doc_title: str, content_type: str = "application/pdf"):
    """
    Upload a local file (PDF or the Backbone Plan JSON handoff) to Firebase Storage
    and return a PERMANENT, browser-fetchable download URL.

    Unchanged from v2.0: a gs:// URI can only be resolved by the Firebase SDK
    (getBytes()/getDownloadURL()), not a plain fetch()/<a href> in the frontend.
    Signed URLs were also considered but expire after 7 days. The fix is a Firebase
    Storage download token attached to the blob's metadata, producing a URL of the
    form:
      https://firebasestorage.googleapis.com/v0/b/<bucket>/o/<path>?alt=media&token=<uuid>
    This URL never expires, works with uniform bucket-level access (no ACL changes
    needed), and is fetchable directly from the browser.

    Returns (permanent_url, storage_path). Falls back to (local_path, None) if
    Firebase isn't configured (so /download keeps working for that session only).
    """
    if not init_firebase():
        return local_path, None

    try:
        filename    = os.path.basename(local_path)
        destination = f"generated-docs/{(project_name or 'project').replace(' ', '_')}/{uuid.uuid4().hex}_{filename}"

        bucket_obj = fb_storage.bucket()
        blob       = bucket_obj.blob(destination)
        blob.upload_from_filename(local_path, content_type=content_type)

        # Attach a Firebase download token → permanent, browser-accessible URL —
        # works fine even with uniform bucket-level access enabled.
        download_token = str(uuid.uuid4())
        blob.metadata  = {"firebaseStorageDownloadTokens": download_token}
        blob.patch()  # persist the metadata so the token URL resolves

        encoded_path  = urllib.parse.quote(destination, safe="")
        permanent_url = (
            f"https://firebasestorage.googleapis.com/v0/b/"
            f"{bucket_obj.name}/o/{encoded_path}"
            f"?alt=media&token={download_token}"
        )

        print(f"✅ Uploaded {filename} → permanent URL generated")
        return permanent_url, destination   # ← real https:// URL + storage path

    except Exception as e:
        print(f"⚠ Upload failed for {local_path}: {e}")
        return local_path, None


# ── App setup ───────────────────────────────────────────────────
app = Flask(__name__)

# ── CORS: explicitly allow only the known frontend origins ──────
CORS(app, resources={r"/api/*": {
    "origins": [
        "https://aihyperdox.com",
        "https://ai-hyperdox.vercel.app",
        "http://localhost:5173",
    ],
    "methods": ["GET", "POST", "OPTIONS"],
    "allow_headers": ["Content-Type", "Authorization"],
    "supports_credentials": False,
}})


# ── Health server (port 8081) ───────────────────────────────────
class HealthHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"OK")
    def log_message(self, format, *args):
        pass

threading.Thread(target=lambda: HTTPServer(("0.0.0.0", 8081), HealthHandler).serve_forever(), daemon=True).start()


# ─────────────────────────────────────────────────────────────────────────────
# Config
# ─────────────────────────────────────────────────────────────────────────────

MODEL = "gpt-4o"
TEMPERATURE = 0.3
MAX_TOKENS = 4096
MAX_RETRIES = 3
RETRY_BACKOFF_SECONDS = 2

MAX_FIELD_CHARS = 4000
MAX_UPLOAD_CHARS = 6000
MAX_HANDOFF_CHARS = 6000
MAX_UPLOAD_FILE_BYTES = 5 * 1024 * 1024  # server-side backstop on raw file bytes (5MB) -
                                          # the frontend caps at 2MB/file client-side, but a
                                          # client limit alone isn't trustworthy; enforce here too

RUN_DIR = "/tmp/aipm_runs"
os.makedirs(RUN_DIR, exist_ok=True)

AGENT_VERSION = "2.2"

# ─────────────────────────────────────────────────────────────────────────────
# System prompt - unchanged agent logic
# ─────────────────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """You are AIPM, an AI project-management advisor built from Monte Turner's
"Be an AI PM" series (Caveman Productions Media). You are the PLANNING agent - the second
of a 3-agent pipeline (Goals & Scope -> Planning -> Execution). You take the governed scope
from the Goals & Scope agent (when provided) plus the user's planning inputs, and produce an
executable plan. Your output is consumed downstream by the Execution agent, so structure and
traceability matter as much as the plan itself.

Your posture is senior enterprise PM: skeptical of vague input, allergic to filler language,
precise about what is knowable versus assumed.

HARD RULES - violating these defeats the purpose of this agent:

1. GROUNDING. Every task, date, resource, and cost figure must trace to something actually
   provided - either the Goals & Scope handoff (if present) or the user's manual planning
   inputs. Never invent team member names, dollar figures, or durations that were not given
   to you or that cannot be reasonably derived from what was given.

2. NO PLACEHOLDER FILLER. If information needed for a section is missing or vague, do not
   invent plausible-sounding filler ("Team A", "TBD", generic 2-week estimates with no
   basis). Instead:
   - Add a specific entry to `open_questions` describing exactly what's missing and why
     it blocks a real answer.
   - Mark the affected field as unresolved, e.g. "start_date": "UNSPECIFIED - see open
     questions", rather than guessing.

3. CROSS-REFERENCING IS MANDATORY.
   - Every row in the Project Timeline must reference a real "wbs_id" from the WBS you
     just generated.
   - Every row in the Resource Allocation Plan must reference a real "wbs_id" or phase
     name from the WBS/timeline, not float independently.
   - The Cost Management Plan's cost baseline categories should map to actual WBS phases,
     not a generic cost-category template.
   - If a Goals & Scope handoff was provided, milestones/risks from it should visibly
     shape the WBS and Timeline (e.g. milestone target dates should show up as phase-gate
     or sprint-end dates), not be ignored in favor of a fresh generic plan.

4. BE CRITICAL. If the stated timeline is unrealistic given the resource list or budget,
   or if a Goals & Scope risk was flagged and nothing in the plan addresses it, say so in
   an open question rather than silently proceeding.

5. SPECIFICITY OVER POLISH. A WBS task like "Anshika: rough-in Episode 1 backgrounds,
   Scenes 1-4, due Jun 25" is correct. A task like "Design team completes visual assets" is
   the generic failure mode to avoid.

6. OUTPUT FORMAT. Respond with ONLY a single valid JSON object matching the schema given
   in the user message. No prose before or after. No markdown code fences.
"""

JSON_SCHEMA_INSTRUCTIONS = """Return a single JSON object with exactly this shape:

{
  "wbs": [
    {"wbs_id": str, "task_name": str, "description": str, "owner": str,
     "est_duration": str, "dependencies": str, "phase_or_epic": str}
  ],
  "timeline": {
    "gantt_summary": str,
    "rows": [
      {"phase_or_sprint": str, "wbs_id": str, "task": str, "start_date": str,
       "end_date": str, "duration": str, "milestone": str, "owner": str, "status": str}
    ]
  },
  "resource_allocation": {
    "rows": [
      {"resource_name_or_role": str, "type": "Human|Tool|Budget", "wbs_id_or_phase": str,
       "allocation_pct": number, "est_hours": number, "cost_rate": str, "notes": str}
    ],
    "utilization_summary": str,
    "overallocation_warnings": [str]
  },
  "cost_plan": {
    "cost_baseline": [
      {"category": str, "wbs_id_or_phase": str, "est_cost": str, "actual": str, "variance": str}
    ],
    "estimation_method": str,
    "contingency_pct": str,
    "control_thresholds": str,
    "reporting_cadence": str,
    "performance_metrics": str
  },
  "open_questions": [str]
}

Rules for this JSON:
- "timeline.rows[].wbs_id" and "resource_allocation.rows[].wbs_id_or_phase" must reference
  real IDs/phases from the "wbs" array you generate in this same response.
- "wbs[].wbs_id" should follow the pattern "1.0", "1.1", "1.1.1" (Waterfall) or
  "EPIC-1", "STORY-1.1" (Agile), matching whichever methodology was specified.
- If you have nothing genuine for open_questions, return an empty array.
"""

# ─────────────────────────────────────────────────────────────────────────────
# File extraction
# ─────────────────────────────────────────────────────────────────────────────

def extract_text_from_file(uploaded_file):
    # BUGFIX (v2.1, unchanged here): a plain string filepath has no .name attribute,
    # so the old getattr(uploaded_file, "name", "uploaded_file") always fell back to
    # the literal string "uploaded_file" for string paths - which is exactly what
    # Gradio's file_count="multiple" hands us per item.
    #
    # NEW in this pass: also accepts {"name": str, "data": str} dicts - what the React
    # frontend's /api/predict call sends. A raw browser File object can't survive
    # JSON.stringify(), so the frontend reads it via FileReader.readAsDataURL() and
    # sends the base64 result (optionally still carrying its "data:<mime>;base64,"
    # prefix, stripped below) instead of a filesystem path. Same pattern as Agent 1's
    # extract_text_from_file().
    is_dict = isinstance(uploaded_file, dict)
    if is_dict:
        name = str(uploaded_file.get("name", "uploaded_file")).lower()
    elif isinstance(uploaded_file, str):
        name = uploaded_file.lower()
    else:
        name = str(getattr(uploaded_file, "name", "uploaded_file")).lower()

    try:
        if is_dict:
            b64 = uploaded_file.get("data") or uploaded_file.get("content") or ""
            if isinstance(b64, str) and b64.strip().lower().startswith("data:") and "," in b64:
                b64 = b64.split(",", 1)[1]   # strip "data:application/pdf;base64," prefix
            data = base64.b64decode(b64) if b64 else b""
        elif hasattr(uploaded_file, "data"):
            data = uploaded_file.data
        elif isinstance(uploaded_file, bytes):
            data = uploaded_file
        elif isinstance(uploaded_file, str):
            with open(uploaded_file, "rb") as f:
                data = f.read()
        else:
            with open(getattr(uploaded_file, "name", uploaded_file), "rb") as f:
                data = f.read()
    except Exception:
        data = b""

    if len(data) > MAX_UPLOAD_FILE_BYTES:
        return f"(Skipped {os.path.basename(name)}: file exceeds {MAX_UPLOAD_FILE_BYTES // (1024 * 1024)}MB server-side limit)"

    text = ""
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(BytesIO(data))
            for page in reader.pages:
                text += page.extract_text() or ""
        elif name.endswith(".docx"):
            doc = docx.Document(BytesIO(data))
            for p in doc.paragraphs:
                text += p.text + "\n"
        elif name.endswith((".txt", ".csv")):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith((".xlsx", ".xls")):
            if XLSX_AVAILABLE:
                wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
                for sheet in wb.worksheets:
                    text += f"\n[Sheet: {sheet.title}]\n"
                    for row in sheet.iter_rows(values_only=True):
                        row_text = "\t".join(str(c) if c is not None else "" for c in row)
                        if row_text.strip():
                            text += row_text + "\n"
            else:
                text = "(openpyxl not installed - cannot read Excel file)"
        else:
            text = f"(Unsupported file type: {name})"
    except Exception as e:
        text = f"(Error reading {name}: {e})"
    return text.strip()


def extract_handoff_json(uploaded_file):
    """Reads a Goals & Scope handoff .json file. Returns (dict_or_None, warning_or_None).

    NEW in this pass: also accepts a {"name": str, "data": str} base64 dict (what the
    React frontend sends), in addition to the existing Gradio-native string path/object
    shapes - same reasoning as extract_text_from_file() above.
    """
    if uploaded_file is None:
        return None, None
    try:
        if isinstance(uploaded_file, dict):
            b64 = uploaded_file.get("data") or uploaded_file.get("content") or ""
            if isinstance(b64, str) and b64.strip().lower().startswith("data:") and "," in b64:
                b64 = b64.split(",", 1)[1]
            raw_bytes = base64.b64decode(b64) if b64 else b""
            if len(raw_bytes) > MAX_UPLOAD_FILE_BYTES:
                return None, f"⚠ Handoff file exceeds {MAX_UPLOAD_FILE_BYTES // (1024 * 1024)}MB server-side limit - ignoring it."
            raw = raw_bytes.decode("utf-8", errors="ignore")
        else:
            path = uploaded_file if isinstance(uploaded_file, str) else getattr(uploaded_file, "name", None)
            with open(path, "r") as f:
                raw = f.read()
        data = json.loads(raw)
        if data.get("handoff_type") != "goals_and_scope_output":
            return None, "⚠ Uploaded file doesn't look like a Goals & Scope handoff package - ignoring it."
        return data, None
    except Exception as e:
        return None, f"⚠ Could not read handoff file: {e}"


def truncate(text, limit):
    if text is None:
        return ""
    text = str(text)
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n...[truncated at {limit} chars]"


# ─────────────────────────────────────────────────────────────────────────────
# LLM call - unchanged agent logic
# ─────────────────────────────────────────────────────────────────────────────

def call_llm(system_msg, user_prompt):
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError(
            "OPENAI_API_KEY is not set. Go to Space Settings -> Repository Secrets "
            "and add OPENAI_API_KEY."
        )
    client = OpenAI(api_key=api_key)

    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            completion = client.chat.completions.create(
                model=MODEL,
                temperature=TEMPERATURE,
                max_tokens=MAX_TOKENS,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": system_msg},
                    {"role": "user", "content": user_prompt},
                ],
            )
            raw = completion.choices[0].message.content
            return json.loads(raw)
        except json.JSONDecodeError as e:
            last_error = f"Model returned invalid JSON on attempt {attempt}: {e}"
        except Exception as e:
            last_error = f"API error on attempt {attempt}: {e}"
        if attempt < MAX_RETRIES:
            time.sleep(RETRY_BACKOFF_SECONDS * attempt)
    raise RuntimeError(f"LLM call failed after {MAX_RETRIES} attempts. Last error: {last_error}")


# ─────────────────────────────────────────────────────────────────────────────
# Coverage validation - unchanged agent logic (v2.2, mirrors Agent 1's validate_coverage())
# ─────────────────────────────────────────────────────────────────────────────

def validate_coverage(data):
    """Flag WBS phases with no corresponding Timeline, Resource Allocation, or
    Cost Baseline entry, instead of silently letting the model under-cover them.

    Checked at the phase/epic level rather than per individual WBS task, since
    Timeline/Resource/Cost are naturally phase-level artifacts in this schema (a
    Resource Allocation row typically covers a whole EPIC, not every STORY inside
    it - matching real output seen from this agent). A phase counts as covered if
    ANY of its WBS task IDs, or the phase name itself, appears in the relevant
    section. Deliberately a crude string-match check, not semantic matching - same
    tradeoff as Agent 1's version: false positives just add an extra
    open_questions entry for a human to dismiss, false negatives are the exact
    silent-omission bug this exists to catch.
    """
    wbs_items = data.get("wbs", [])
    if not wbs_items:
        return data

    wbs_ids_by_phase = {}
    for w in wbs_items:
        phase = str(w.get("phase_or_epic", "")).strip()
        if not phase:
            continue
        wbs_ids_by_phase.setdefault(phase, set()).add(str(w.get("wbs_id", "")).strip())

    timeline_rows = data.get("timeline", {}).get("rows", [])
    resource_rows = data.get("resource_allocation", {}).get("rows", [])
    cost_rows = data.get("cost_plan", {}).get("cost_baseline", [])

    timeline_refs = ({str(r.get("wbs_id", "")).strip() for r in timeline_rows} |
                      {str(r.get("phase_or_sprint", "")).strip() for r in timeline_rows})
    resource_refs = {str(r.get("wbs_id_or_phase", "")).strip() for r in resource_rows}
    cost_refs = {str(r.get("wbs_id_or_phase", "")).strip() for r in cost_rows}

    open_qs = data.setdefault("open_questions", [])
    for phase, ids in wbs_ids_by_phase.items():
        covered_by = ids | {phase}
        if not (covered_by & timeline_refs):
            open_qs.append(
                f"No Timeline entry appears tied to WBS phase '{phase}' - "
                f"confirm scheduling for this phase."
            )
        if not (covered_by & resource_refs):
            open_qs.append(
                f"No Resource Allocation entry appears tied to WBS phase '{phase}' - "
                f"confirm ownership/staffing for this phase."
            )
        if not (covered_by & cost_refs):
            open_qs.append(
                f"No Cost Baseline entry appears tied to WBS phase '{phase}' - "
                f"confirm a cost estimate exists for this phase."
            )
    return data


def validate_unspecified_markers(data):
    """If any field was marked UNSPECIFIED by the model but open_questions came
    back empty, flag that explicitly rather than letting a plan look complete
    while actually being silently hollow.

    This is the direct fix for a real observed pattern: a Cost Management Plan
    where every single field (est_cost, actual, variance, estimation_method,
    contingency, thresholds, cadence, metrics) came back "UNSPECIFIED - see open
    questions" with no corresponding open_questions entries explaining why -
    exactly the failure mode the anti-genericism rule is supposed to prevent.
    """
    marker = "UNSPECIFIED - see open questions"
    found = {"flag": False}

    def scan(obj):
        if found["flag"]:
            return
        if isinstance(obj, dict):
            for v in obj.values():
                scan(v)
        elif isinstance(obj, list):
            for v in obj:
                scan(v)
        elif isinstance(obj, str) and marker in obj:
            found["flag"] = True

    scan(data.get("wbs", []))
    scan(data.get("timeline", {}))
    scan(data.get("resource_allocation", {}))
    scan(data.get("cost_plan", {}))

    open_qs = data.setdefault("open_questions", [])
    if found["flag"] and not open_qs:
        open_qs.append(
            "One or more plan fields were marked UNSPECIFIED but no open_questions "
            "were logged explaining why - review the plan for unresolved gaps "
            "before treating it as complete."
        )
    return data


# ─────────────────────────────────────────────────────────────────────────────
# PDF rendering - unchanged agent logic
# ─────────────────────────────────────────────────────────────────────────────

def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("AIPMTitle", parent=styles["Title"],
                               textColor=colors.HexColor("#1a3c6e"), fontSize=18))
    styles.add(ParagraphStyle("AIPMH2", parent=styles["Heading2"],
                               textColor=colors.HexColor("#2563eb"), fontSize=13,
                               spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle("AIPMCell", parent=styles["Normal"], fontSize=8.5, leading=11))
    return styles


def _table_from_dicts(rows, columns, styles):
    header = [Paragraph(f"<b>{label}</b>", styles["AIPMCell"]) for _, label, _ in columns]
    data = [header]
    for row in rows:
        data.append([Paragraph(str(row.get(key, "")), styles["AIPMCell"]) for key, _, _ in columns])
    total_width = 7.0 * inch
    col_widths = [total_width * frac for _, _, frac in columns]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3c6e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6fb")]),
    ]))
    return table


def _bullet_list(items, styles):
    return ListFlowable(
        [ListItem(Paragraph(str(i), styles["Normal"])) for i in items] or
        [ListItem(Paragraph("(none)", styles["Normal"]))],
        bulletType="bullet",
    )


def build_all_pdfs(run_dir, project_name, data):
    styles = _styles()
    safe_pn = (project_name or "Project").replace(" ", "_")
    pdfs = {}

    def new_doc(title):
        path = os.path.join(run_dir, f"{safe_pn}_{title.replace(' ', '_')}.pdf")
        doc = SimpleDocTemplate(path, pagesize=LETTER,
                                 leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                                 topMargin=0.75 * inch, bottomMargin=0.75 * inch)
        story = [
            Paragraph(project_name or "Project", styles["AIPMTitle"]),
            Spacer(1, 6),
            Paragraph(title, styles["AIPMH2"]),
            Spacer(1, 10),
        ]
        return path, doc, story

    # WBS
    path, doc, story = new_doc("Work Breakdown Structure (WBS)")
    story.append(_table_from_dicts(
        data.get("wbs", []),
        [("wbs_id", "WBS ID", 0.09), ("task_name", "Task", 0.18), ("description", "Description", 0.27),
         ("owner", "Owner", 0.12), ("est_duration", "Duration", 0.10), ("dependencies", "Deps", 0.12),
         ("phase_or_epic", "Phase/Epic", 0.12)],
        styles,
    ))
    doc.build(story)
    pdfs["Work Breakdown Structure (WBS)"] = path

    # Timeline
    path, doc, story = new_doc("Project Timeline")
    timeline = data.get("timeline", {})
    story.append(Paragraph(f"<i>{timeline.get('gantt_summary', '')}</i>", styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(_table_from_dicts(
        timeline.get("rows", []),
        [("phase_or_sprint", "Phase/Sprint", 0.13), ("wbs_id", "WBS ID", 0.08), ("task", "Task", 0.19),
         ("start_date", "Start", 0.10), ("end_date", "End", 0.10), ("duration", "Dur.", 0.08),
         ("milestone", "Milestone", 0.12), ("owner", "Owner", 0.10), ("status", "Status", 0.10)],
        styles,
    ))
    doc.build(story)
    pdfs["Project Timeline"] = path

    # Resource Allocation
    path, doc, story = new_doc("Resource Allocation Plan")
    ra = data.get("resource_allocation", {})
    story.append(_table_from_dicts(
        ra.get("rows", []),
        [("resource_name_or_role", "Resource", 0.18), ("type", "Type", 0.10),
         ("wbs_id_or_phase", "WBS/Phase", 0.14), ("allocation_pct", "Alloc %", 0.09),
         ("est_hours", "Hours", 0.09), ("cost_rate", "Rate", 0.12), ("notes", "Notes", 0.28)],
        styles,
    ))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Utilization Summary", styles["AIPMH2"]))
    story.append(Paragraph(ra.get("utilization_summary", ""), styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Over-Allocation Warnings", styles["AIPMH2"]))
    story.append(_bullet_list(ra.get("overallocation_warnings", []), styles))
    doc.build(story)
    pdfs["Resource Allocation Plan"] = path

    # Cost Management Plan
    path, doc, story = new_doc("Cost Management Plan")
    cp = data.get("cost_plan", {})
    story.append(Paragraph("Cost Baseline", styles["AIPMH2"]))
    story.append(_table_from_dicts(
        cp.get("cost_baseline", []),
        [("category", "Category", 0.22), ("wbs_id_or_phase", "WBS/Phase", 0.15), ("est_cost", "Est. Cost", 0.18),
         ("actual", "Actual", 0.18), ("variance", "Variance", 0.27)],
        styles,
    ))
    for label, key in [("Cost Estimation Method", "estimation_method"),
                        ("Budget Contingency", "contingency_pct"),
                        ("Cost Control Thresholds", "control_thresholds"),
                        ("Reporting Cadence", "reporting_cadence"),
                        ("Cost Performance Metrics", "performance_metrics")]:
        story.append(Spacer(1, 8))
        story.append(Paragraph(label, styles["AIPMH2"]))
        story.append(Paragraph(str(cp.get(key, "")), styles["Normal"]))
    doc.build(story)
    pdfs["Cost Management Plan"] = path

    # Combined Project Plan Guardrail Document (all sections + open questions)
    path, doc, story = new_doc("Project Plan Guardrail Document (Combined)")
    story.append(Paragraph("Work Breakdown Structure", styles["AIPMH2"]))
    story.append(_table_from_dicts(
        data.get("wbs", []),
        [("wbs_id", "WBS ID", 0.09), ("task_name", "Task", 0.18), ("description", "Description", 0.27),
         ("owner", "Owner", 0.12), ("est_duration", "Duration", 0.10), ("dependencies", "Deps", 0.12),
         ("phase_or_epic", "Phase/Epic", 0.12)],
        styles,
    ))
    story.append(PageBreak())
    story.append(Paragraph("Project Timeline", styles["AIPMH2"]))
    story.append(Paragraph(f"<i>{timeline.get('gantt_summary', '')}</i>", styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(_table_from_dicts(
        timeline.get("rows", []),
        [("phase_or_sprint", "Phase/Sprint", 0.13), ("wbs_id", "WBS ID", 0.08), ("task", "Task", 0.19),
         ("start_date", "Start", 0.10), ("end_date", "End", 0.10), ("duration", "Dur.", 0.08),
         ("milestone", "Milestone", 0.12), ("owner", "Owner", 0.10), ("status", "Status", 0.10)],
        styles,
    ))
    story.append(PageBreak())
    story.append(Paragraph("Resource Allocation Plan", styles["AIPMH2"]))
    story.append(_table_from_dicts(
        ra.get("rows", []),
        [("resource_name_or_role", "Resource", 0.18), ("type", "Type", 0.10),
         ("wbs_id_or_phase", "WBS/Phase", 0.14), ("allocation_pct", "Alloc %", 0.09),
         ("est_hours", "Hours", 0.09), ("cost_rate", "Rate", 0.12), ("notes", "Notes", 0.28)],
        styles,
    ))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Utilization Summary", styles["Normal"]))
    story.append(Paragraph(ra.get("utilization_summary", ""), styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Over-Allocation Warnings", styles["Normal"]))
    story.append(_bullet_list(ra.get("overallocation_warnings", []), styles))
    story.append(PageBreak())
    story.append(Paragraph("Cost Management Plan", styles["AIPMH2"]))
    story.append(_table_from_dicts(
        cp.get("cost_baseline", []),
        [("category", "Category", 0.22), ("wbs_id_or_phase", "WBS/Phase", 0.15), ("est_cost", "Est. Cost", 0.18),
         ("actual", "Actual", 0.18), ("variance", "Variance", 0.27)],
        styles,
    ))
    for label, key in [("Cost Estimation Method", "estimation_method"),
                        ("Budget Contingency", "contingency_pct"),
                        ("Cost Control Thresholds", "control_thresholds"),
                        ("Reporting Cadence", "reporting_cadence"),
                        ("Cost Performance Metrics", "performance_metrics")]:
        story.append(Spacer(1, 6))
        story.append(Paragraph(label, styles["Normal"]))
        story.append(Paragraph(str(cp.get(key, "")), styles["Normal"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Open Questions", styles["AIPMH2"]))
    story.append(_bullet_list(data.get("open_questions", []), styles))
    doc.build(story)
    pdfs["Project Plan Guardrail Document (Combined)"] = path

    return pdfs


# ─────────────────────────────────────────────────────────────────────────────
# Main generation function
# ─────────────────────────────────────────────────────────────────────────────

def generate_project_plan(project_name, milestones, timeline_input, resources,
                           budget_text_input, supporting_uploads, methodology, handoff_upload):
    if not any([project_name, milestones, timeline_input, resources,
                budget_text_input, supporting_uploads, handoff_upload]):
        return (
            "⚠ Please provide at least a Project Name and some planning details, "
            "or upload a Goals & Scope handoff package.",
            None, None, None, None, None, None,
        )

    run_id = uuid.uuid4().hex[:10]
    run_dir = os.path.join(RUN_DIR, run_id)
    os.makedirs(run_dir, exist_ok=True)

    handoff_data, handoff_warning = extract_handoff_json(handoff_upload)

    # Supporting documents: any file type/count, same pattern as Agent 1 (Goals & Scope).
    # Not limited to budget files - resource lists, timelines, etc. can go here too.
    extra_text = ""
    if supporting_uploads:
        for f in supporting_uploads:
            extracted = extract_text_from_file(f)
            if isinstance(f, dict):
                fname = f.get("name", "uploaded file")
            elif isinstance(f, str):
                fname = f
            else:
                fname = getattr(f, "name", "uploaded file")
            fname = os.path.basename(str(fname))
            extra_text += f"\n\n[Extracted from {fname}]\n{extracted}"

    handoff_context = "Not provided - no Goals & Scope handoff package was uploaded."
    if handoff_data:
        handoff_context = truncate(json.dumps({
            "project_name": handoff_data.get("project_name"),
            "scope": handoff_data.get("scope"),
            "milestones": handoff_data.get("milestones"),
            "risks": handoff_data.get("risks"),
            "resource_teams": handoff_data.get("resource_teams"),
            "open_questions": handoff_data.get("open_questions"),
        }, indent=2), MAX_HANDOFF_CHARS)

    context = f"""Project: {project_name}
Methodology: {methodology}

Goals & Scope Handoff Package (from Agent 1, if provided):
{handoff_context}

Manual Milestones Input (may overlap with or extend the handoff above):
{truncate(milestones, MAX_FIELD_CHARS) or 'Not specified'}

Manual High-Level Timeline Input:
{truncate(timeline_input, MAX_FIELD_CHARS) or 'Not specified'}

Manual Resource List Input:
{truncate(resources, MAX_FIELD_CHARS) or 'Not specified'}

Budget / Financial Data (pasted directly):
{truncate(budget_text_input, MAX_FIELD_CHARS) or 'Not provided'}

Uploaded Supporting Documents Content (any type - budget, resources, timeline, etc.):
{truncate(extra_text, MAX_UPLOAD_CHARS) or 'None uploaded'}

Date: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}
"""

    methodology_note = (
        "Use Agile sprint-based decomposition - organize by Epic -> User Story -> Task with "
        "Sprint assignments, wbs_id like EPIC-1 / STORY-1.1."
        if methodology == "Agile" else
        "Use classic Waterfall hierarchical WBS decomposition aligned to phase gates with "
        "sign-off points, wbs_id like 1.0 / 1.1 / 1.1.1."
    )

    user_prompt = f"{context}\n{methodology_note}\n\n{JSON_SCHEMA_INSTRUCTIONS}"

    try:
        data = call_llm(SYSTEM_PROMPT, user_prompt)
    except Exception as e:
        return (f"⚠ Generation failed: {e}", None, None, None, None, None, None)

    data = validate_coverage(data)
    data = validate_unspecified_markers(data)

    try:
        pdfs = build_all_pdfs(run_dir, project_name, data)
    except Exception as e:
        return (f"⚠ Document rendering failed: {e}", None, None, None, None, None, None)

    # Handoff package for Agent 3 (Execution) - "Backbone Plan"
    backbone = {
        "handoff_type": "planning_output",
        "handoff_version": "1.0",
        "source_agent": f"planning_v{AGENT_VERSION}",
        "upstream_handoff_used": bool(handoff_data),
        "project_name": project_name,
        "methodology": methodology,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "wbs": data.get("wbs", []),
        "timeline": data.get("timeline", {}),
        "resource_allocation": data.get("resource_allocation", {}),
        "cost_plan": data.get("cost_plan", {}),
        "open_questions": data.get("open_questions", []),
    }
    backbone_path = os.path.join(run_dir, f"{(project_name or 'project').replace(' ', '_')}_backbone_plan.json")
    with open(backbone_path, "w") as f:
        json.dump(backbone, f, indent=2)

    # Upload every output (5 PDFs + Backbone Plan JSON) to Firebase Storage → permanent,
    # browser-fetchable URLs that survive server restarts/redeploys.
    uploaded = {}
    for title, local_path in pdfs.items():
        url, storage_path = upload_to_storage(local_path, project_name, title, content_type="application/pdf")
        uploaded[title] = {"url": url, "path": storage_path}
    backbone_url, backbone_storage_path = upload_to_storage(
        backbone_path, project_name, "Backbone Plan", content_type="application/json"
    )
    uploaded_backbone = {"url": backbone_url, "path": backbone_storage_path}

    status_lines = []
    if handoff_warning:
        status_lines.append(handoff_warning)
    status_lines.append(f"✅ All planning documents generated for '{project_name}'" +
                         (" (using Goals & Scope handoff)" if handoff_data else "") + ".")
    if data.get("open_questions"):
        status_lines.append(f"⚠ {len(data['open_questions'])} open question(s) flagged - see the combined document.")

    return (
        "\n".join(status_lines),
        uploaded["Work Breakdown Structure (WBS)"],
        uploaded["Project Timeline"],
        uploaded["Resource Allocation Plan"],
        uploaded["Cost Management Plan"],
        uploaded["Project Plan Guardrail Document (Combined)"],
        uploaded_backbone,
    )


# ─────────────────────────────────────────────────────────────────────────────
# Flask REST API — this is what the aihyperdox.com frontend actually calls
# ─────────────────────────────────────────────────────────────────────────────

def _guess_mimetype(path):
    """Same hardening as Agent 1: the /download fallback used to hardcode
    application/pdf even for the .json Backbone Plan. Only matters when Firebase
    isn't configured and /download is actually serving files."""
    return "application/json" if path.lower().endswith(".json") else "application/pdf"


@app.route("/api/predict", methods=["POST", "OPTIONS"])
def api_predict():
    if request.method == "OPTIONS":
        return "", 204
    try:
        data       = request.json or {}
        input_data = data.get("data", [])

        if len(input_data) < 4:
            return jsonify({"error": "Missing required fields"}), 400

        # Positional order matches generate_project_plan()'s v2.2 signature - NOT the
        # same order as v2.0/v2.1. See module docstring for the full before/after.
        project_name       = input_data[0] or ""
        milestones         = input_data[1] or ""
        timeline_input     = input_data[2] or ""
        resources          = input_data[3] or ""
        budget_text_input  = input_data[4] if len(input_data) > 4 else ""
        supporting_uploads = input_data[5] if len(input_data) > 5 else []
        methodology        = input_data[6] if len(input_data) > 6 else "Agile"
        handoff_upload      = input_data[7] if len(input_data) > 7 else None

        result = generate_project_plan(
            project_name, milestones, timeline_input, resources,
            budget_text_input, supporting_uploads or [], methodology, handoff_upload,
        )
        return jsonify({"data": result}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/download", methods=["GET"])
def download_file():
    """
    Fallback download endpoint — only used when Firebase Storage is not configured
    and files are still on the Railway /tmp filesystem (i.e. same session). Once
    Firebase is configured, generate_project_plan() returns a permanent
    firebasestorage.googleapis.com URL instead, and this endpoint is never hit for
    new runs.
    """
    path = request.args.get("path")
    if not path:
        abort(400, description="Missing 'path' query parameter.")
    if not path.startswith(RUN_DIR):
        abort(403, description="Access denied.")
    if not os.path.exists(path):
        abort(404, description=f"File not found: {path}. The server may have restarted — please regenerate the documents.")
    return send_file(
        path,
        as_attachment=True,
        download_name=os.path.basename(path),
        mimetype=_guess_mimetype(path),
    )


@app.route("/api/health", methods=["GET"])
def health_check():
    return jsonify({"status": "ok", "firebase": _firebase_ready, "agent_version": AGENT_VERSION}), 200


# ─────────────────────────────────────────────────────────────────────────────
# Gradio UI
# ─────────────────────────────────────────────────────────────────────────────

with gr.Blocks(title="AIPM - Project Plan Documents") as demo:

    gr.HTML(
        "<div style='text-align:center'><h1>🗂️ AIPM - Planning Agent</h1>"
        "<p>Agent 2 of 3 · Goals & Scope → Planning → Execution</p></div>"
    )

    with gr.Row():
        with gr.Column(scale=1):
            handoff_upload = gr.File(
                label="🔗 Goals & Scope Handoff Package (optional, .json from Agent 1)",
                file_count="single",
            )
            project_name = gr.Textbox(
                label="Project Name", placeholder="e.g. Uggalot Episode 1 Production", lines=1,
            )
            methodology = gr.Radio(
                choices=["Agile", "Waterfall"], value="Agile", label="📐 Methodology",
                info="Agile = sprint/epic decomposition. Waterfall = phase-gate planning.",
            )
            milestones = gr.Textbox(
                label="Milestones (optional if handoff package provided - will be merged)",
                placeholder="e.g.\nM1 – Script Finalized | 2026-06-20\nM2 – Animation Draft Complete | 2026-07-15",
                lines=5,
            )
            timeline = gr.Textbox(
                label="High-Level Timeline",
                placeholder="e.g.\nPhase 1: Pre-Production (Jun 1 – Jun 30)\nPhase 2: Production (Jul 1 – Jul 31)",
                lines=4,
            )
            resources = gr.Textbox(
                label="Resource List (optional if handoff package provided - will be merged)",
                placeholder="e.g.\nAnshika – Illustrator (full-time)\nDeepanshu – Animator Intern",
                lines=5,
            )
            budget_text = gr.Textbox(
                label="Budget / Financial Data (optional - paste directly)",
                placeholder="e.g.\nAnthropic API usage (Haiku): ~$15-30/mo est.\nDev time: $0 internal, sweat equity",
                lines=4,
            )
            uploads = gr.File(
                label="Upload Supporting Documents (optional, any format - budget, resources, timeline, etc.)",
                file_count="multiple",
            )
            submit_btn = gr.Button("🚀 Generate Project Plan Documents", variant="primary")

        with gr.Column(scale=1):
            status = gr.Textbox(label="Status", interactive=False, lines=4)
            pdf_wbs = gr.File(label="📥 Work Breakdown Structure (WBS)")
            pdf_tl = gr.File(label="📥 Project Timeline")
            pdf_ra = gr.File(label="📥 Resource Allocation Plan")
            pdf_cm = gr.File(label="📥 Cost Management Plan")
            pdf_combined = gr.File(label="📥 Project Plan Guardrail Document (Combined)")
            backbone_file = gr.File(label="🔗 Backbone Plan - Handoff for Execution Agent (JSON)")

    submit_btn.click(
        fn=generate_project_plan,
        inputs=[project_name, milestones, timeline, resources, budget_text, uploads,
                methodology, handoff_upload],
        outputs=[status, pdf_wbs, pdf_tl, pdf_ra, pdf_cm, pdf_combined, backbone_file],
        show_progress=True,
    )

    gr.HTML(
        f"<p style='text-align:center;color:#bbb;font-size:11px;margin-top:16px;'>"
        f"© 2026 Caveman Productions Media &nbsp;·&nbsp; AIPM Planning Agent v{AGENT_VERSION} "
        f"&nbsp;·&nbsp; Agent 2 of 3</p>"
    )

if __name__ == "__main__":
    def run_gradio():
        demo.launch(
            server_name="127.0.0.1",
            server_port=7860,
            share=False,
            quiet=True,
            theme=gr.themes.Soft(primary_hue="blue", neutral_hue="gray"),
        )

    threading.Thread(target=run_gradio, daemon=True).start()

    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port, debug=False)