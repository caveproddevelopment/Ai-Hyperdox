from openai import OpenAI
from reportlab.lib.pagesizes import LETTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from PyPDF2 import PdfReader
from io import BytesIO
import docx
import gradio as gr
import os
from datetime import datetime, timezone

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False


# ── File extraction ────────────────────────────────────────────────────────────

def extract_text_from_file(uploaded_file):
    name = getattr(uploaded_file, "name", "uploaded_file").lower()
    try:
        with open(uploaded_file, "rb") as f:
            data = f.read()
    except Exception:
        data = b""

    text = ""
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(BytesIO(data))
            for page in reader.pages:
                text += page.extract_text() or ""
        elif name.endswith(".docx"):
            doc = docx.Document(BytesIO(data))
            for p in doc.paragraphs:
                text += p.text + "\n"
        elif name.endswith(".txt"):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(".csv"):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(".xlsx") or name.endswith(".xls"):
            if XLSX_AVAILABLE:
                wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
                for sheet in wb.worksheets:
                    text += f"\n[Sheet: {sheet.title}]\n"
                    for row in sheet.iter_rows(values_only=True):
                        row_text = "\t".join(
                            str(cell) if cell is not None else "" for cell in row
                        )
                        if row_text.strip():
                            text += row_text + "\n"
            else:
                text = "(openpyxl not installed — cannot read Excel file)"
        else:
            text = f"(Unsupported file type: {name})"
    except Exception as e:
        text = f"(Error reading {name}: {e})"
    return text.strip()


# ── PDF builder ────────────────────────────────────────────────────────────────

def create_pdf(project_name, doc_title, content):
    styles   = getSampleStyleSheet()
    safe_pn  = (project_name or "Project").replace(" ", "_")
    safe_dt  = doc_title.replace(" ", "_")
    filepath = f"/tmp/{safe_pn}_{safe_dt}.pdf"

    title_style = ParagraphStyle(
        "PPTitle", parent=styles["Title"],
        textColor=colors.HexColor("#1a3c6e"), fontSize=18
    )
    h2_style = ParagraphStyle(
        "PPH2", parent=styles["Heading2"],
        textColor=colors.HexColor("#2563eb"), fontSize=13
    )

    doc = SimpleDocTemplate(
        filepath, pagesize=LETTER,
        leftMargin=0.75*inch, rightMargin=0.75*inch,
        topMargin=0.75*inch,  bottomMargin=0.75*inch,
    )
    story = [
        Paragraph(safe_pn.replace("_", " "), title_style),
        Spacer(1, 6),
        Paragraph(doc_title, h2_style),
        Spacer(1, 12),
        Paragraph(content.replace("\n", "<br/>"), styles["Normal"]),
    ]
    doc.build(story)
    return filepath


# ── GPT call (lazy client — no module-level init) ──────────────────────────────

def call_gpt(system_msg, user_prompt):
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError(
            "OPENAI_API_KEY is not set. "
            "Go to Space Settings → Repository Secrets and add OPENAI_API_KEY."
        )
    _client = OpenAI(api_key=api_key)
    completion = _client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": system_msg},
            {"role": "user",   "content": user_prompt},
        ],
    )
    return completion.choices[0].message.content


# ── Main generation function ───────────────────────────────────────────────────

def generate_project_plan(project_name, milestones, timeline, resources, budget_file, methodology):
    if not any([project_name, milestones, timeline, resources]):
        return (
            "⚠ Please provide at least a Project Name and some planning details.",
            None, None, None, None,
        )

    budget_text = ""
    if budget_file is not None:
        budget_text = extract_text_from_file(budget_file)

    context = f"""Project: {project_name}
Methodology: {methodology}

Milestones:
{milestones or 'Not specified'}

High-Level Timeline:
{timeline or 'Not specified'}

Resource List:
{resources or 'Not specified'}

Budget / Financial Data:
{budget_text[:4000] if budget_text else 'Not provided'}

Date: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}
"""

    system_msg = (
        "You are AIPM, an expert AI Project Manager from Monte Turner's 'Be an AI PM' series. "
        f"Apply {methodology} methodology throughout all outputs. "
        "Generate concise, professional, structured project planning documentation."
    )

    methodology_note = (
        "Use Agile sprint-based decomposition — organize by Epic → User Story → Task with Sprint assignments."
        if methodology == "Agile"
        else "Use classic Waterfall hierarchical WBS decomposition aligned to phase gates with sign-off points."
    )

    prompts = {
        "Work Breakdown Structure (WBS)": (
            f"{context}\n{methodology_note}\n\n"
            "Generate a complete Work Breakdown Structure (WBS).\n"
            "Structure: Level 1 = Major Phases, Level 2 = Deliverables, Level 3 = Work Packages/Tasks.\n"
            "Include WBS ID codes (1.0, 1.1, 1.1.1...).\n"
            "Format as a table: WBS ID | Task Name | Description | Owner | Est. Duration | Dependencies."
        ),
        "Project Timeline": (
            f"{context}\n\n"
            f"Create a detailed project timeline aligned to {methodology} principles.\n"
            "Include a Gantt-style text summary at the top.\n"
            "Then a table: Phase/Sprint | Task | Start Date | End Date | Duration | Milestone | Owner | Status.\n"
            "Agile: organize by Sprint/Iteration. Waterfall: organize by phase gates."
        ),
        "Resource Allocation Plan": (
            f"{context}\n\n"
            "Develop a Resource Allocation Plan.\n"
            "Table: Resource Name/Role | Type (Human/Tool/Budget) | Assigned Phase or Sprint | "
            "Allocation % | Est. Hours | Cost Rate (if known) | Notes.\n"
            "Include a resource utilization summary and flag any over-allocations."
        ),
        "Cost Management Plan": (
            f"{context}\n\n"
            "Develop a Cost Management Plan with these sections:\n"
            "1. Cost Baseline — itemized by phase or work package (table: Category | Est. Cost | Actual | Variance)\n"
            "2. Cost Estimation Method — how costs were derived\n"
            "3. Budget Contingency — recommended reserve %\n"
            "4. Cost Control Thresholds — variance triggers for escalation\n"
            "5. Reporting Cadence — who reviews costs and how often\n"
            "6. Cost Performance Metrics — CPI, SPI where applicable"
        ),
    }

    pdfs   = {}
    errors = []
    for title, prompt in prompts.items():
        try:
            content   = call_gpt(system_msg, prompt)
            pdfs[title] = create_pdf(project_name, title, content)
        except Exception as e:
            errors.append(f"{title}: {e}")
            pdfs[title] = None

    if errors:
        status = f"⚠ Errors occurred:\n" + "\n".join(errors)
    else:
        status = f"✅ All four Project Plan documents generated for '{project_name}'!"

    return (
        status,
        pdfs.get("Work Breakdown Structure (WBS)"),
        pdfs.get("Project Timeline"),
        pdfs.get("Resource Allocation Plan"),
        pdfs.get("Cost Management Plan"),
    )


# ── Gradio UI ──────────────────────────────────────────────────────────────────

with gr.Blocks(
    title="AIPM – Project Plan Documents",
    theme=gr.themes.Soft(primary_hue="blue", neutral_hue="gray"),
    css="""
        .header { text-align:center; padding:14px 0 6px; }
        .header h1 { font-size:1.75rem; font-weight:700; color:#1a3c6e; }
        .header p  { color:#555; font-size:.93rem; margin-top:3px; }
        footer { visibility:hidden; }
    """,
) as demo:

    gr.HTML("""
        <div class="header">
            <h1>🗂️ AIPM – Project Plan Documents</h1>
            <p>Monte Turner's <em>Be an AI PM</em> series &nbsp;·&nbsp;
               WBS &nbsp;·&nbsp; Timeline &nbsp;·&nbsp;
               Resource Allocation &nbsp;·&nbsp; Cost Management Plan</p>
        </div>
    """)

    with gr.Row():

        # ── Left column: inputs ────────────────────────────────────────────
        with gr.Column(scale=1):

            project_name = gr.Textbox(
                label="Project Name",
                placeholder="e.g. Uggalot Episode 1 Production",
                lines=1,
            )

            methodology = gr.Radio(
                choices=["Agile", "Waterfall"],
                value="Agile",
                label="📐 Methodology",
                info="Agile = sprint/epic decomposition.  Waterfall = phase-gate planning.",
            )

            milestones = gr.Textbox(
                label="Milestones",
                placeholder=(
                    "e.g.\n"
                    "M1 – Script Finalized | 2026-06-20\n"
                    "M2 – Animation Draft Complete | 2026-07-15\n"
                    "M3 – Episode 1 Launch | 2026-08-01"
                ),
                lines=6,
            )

            timeline = gr.Textbox(
                label="High-Level Timeline",
                placeholder=(
                    "e.g.\n"
                    "Phase 1: Pre-Production  (Jun 1 – Jun 30)\n"
                    "Phase 2: Production      (Jul 1 – Jul 31)\n"
                    "Phase 3: Post-Production (Aug 1 – Aug 15)\n"
                    "Phase 4: Launch          (Aug 16+)"
                ),
                lines=5,
            )

            resources = gr.Textbox(
                label="Resource List",
                placeholder=(
                    "e.g.\n"
                    "Anshika – Illustrator (full-time)\n"
                    "Helna   – Illustrator (full-time)\n"
                    "Deepanshu – Animator Intern\n"
                    "Ratim   – Animator Intern\n"
                    "Aatish  – App Developer Intern"
                ),
                lines=6,
            )

            budget_file = gr.File(
                label="📊 Budget Spreadsheet — optional (XLSX / CSV / PDF / TXT)",
                file_count="single",
            )

            submit_btn = gr.Button("🚀 Generate Project Plan Documents", variant="primary")

        # ── Right column: outputs ──────────────────────────────────────────
        with gr.Column(scale=1):
            status  = gr.Textbox(label="Status", interactive=False, lines=4)
            pdf_wbs = gr.File(label="📥 Work Breakdown Structure (WBS)")
            pdf_tl  = gr.File(label="📥 Project Timeline")
            pdf_ra  = gr.File(label="📥 Resource Allocation Plan")
            pdf_cm  = gr.File(label="📥 Cost Management Plan")

    submit_btn.click(
        fn=generate_project_plan,
        inputs=[project_name, milestones, timeline, resources, budget_file, methodology],
        outputs=[status, pdf_wbs, pdf_tl, pdf_ra, pdf_cm],
        show_progress=True,
    )

    gr.HTML(
        "<p style='text-align:center;color:#bbb;font-size:11px;margin-top:16px;'>"
        "© 2026 Caveman Productions Media &nbsp;·&nbsp; AIPM Project Plan Agent &nbsp;·&nbsp; Agent 2 of 5</p>"
    )

if __name__ == "__main__":
    demo.launch()
